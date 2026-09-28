"""受発注サンプルファイル (PDF x3 / Excel x3) を生成するスクリプト。

使い方:
    pip install -r tools/requirements.txt
    python tools/generate_samples.py

生成先:
    samples/pdf/*.pdf
    samples/excel/*.xlsx

3 つの PDF / 3 つの Excel はそれぞれ異なるフォーマット (項目の並び順、
ラベルの表記ゆれ、レイアウト) ですが、Web アプリへ転記する項目
(注文番号 / 受注日 / 取引先 / 納品希望日 / 商品コード / 数量 / 単価 /
通貨 / 支払条件 / 出荷方法 / 担当者名 / 備考) はすべて含みます。
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF_DIR = os.path.join(ROOT, "samples", "pdf")
EXCEL_DIR = os.path.join(ROOT, "samples", "excel")

FONT_NAME = "HeiseiKakuGo-W5"

PDF_ORDERS = [
    {
        "file": "order-001",
        "orderNumber": "PO-2026-0001",
        "orderDate": "2026/04/06",
        "customerCode": "C-1001",
        "customerName": "山田製作所",
        "deliveryDate": "2026/05/11",
        "productCode": "P-2001",
        "productName": "精密ボルト M6",
        "quantity": 1200,
        "unitPrice": 85,
        "currency": "JPY",
        "paymentTerms": "30日後支払",
        "shippingMethod": "通常便",
        "picName": "鈴木 一郎",
        "inspection": "必要",
        "remarks": "初回ロットのため検査成績書を添付のこと。",
    },
    {
        "file": "order-002",
        "orderNumber": "PO-2026-0042",
        "orderDate": "2026/04/14",
        "customerCode": "C-1003",
        "customerName": "北斗マテリアル",
        "deliveryDate": "2026/04/28",
        "productCode": "P-2002",
        "productName": "ステンレス板 SUS304",
        "quantity": 45,
        "unitPrice": 12800,
        "currency": "JPY",
        "paymentTerms": "月末締め翌月末払い",
        "shippingMethod": "急送便",
        "picName": "佐藤 花子",
        "inspection": "不要",
        "remarks": "工場直送。荷受は第2倉庫。",
    },
    {
        "file": "order-003",
        "orderNumber": "PO-2026-0117",
        "orderDate": "2026/05/07",
        "customerCode": "C-1005",
        "customerName": "九州テクノサービス",
        "deliveryDate": "2026/06/19",
        "productCode": "P-2005",
        "productName": "制御基板 Rev.2",
        "quantity": 30,
        "unitPrice": 24500,
        "currency": "JPY",
        "paymentTerms": "前払い",
        "shippingMethod": "引取",
        "picName": "田中 健",
        "inspection": "必要",
        "remarks": "納品前に動作試験報告書を提出すること。",
    },
]


def _styles():
    return {
        "title": ParagraphStyle("title", fontName=FONT_NAME, fontSize=16, leading=22),
        "normal": ParagraphStyle("normal", fontName=FONT_NAME, fontSize=10, leading=16),
        "small": ParagraphStyle("small", fontName=FONT_NAME, fontSize=8, leading=12),
    }


def _table(data, col_widths, header_rows=0):
    table = Table(data, colWidths=col_widths)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8c99a8")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header_rows:
        style.append(("BACKGROUND", (0, 0), (-1, header_rows - 1), colors.HexColor("#dde5ee")))
    table.setStyle(TableStyle(style))
    return table


def build_pdf_format_a(order, path):
    """フォーマット A: 一般的な 2 列の注文書 (項目名 / 値)。"""
    styles = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    rows = [
        ["注文番号", order["orderNumber"], "注文日", order["orderDate"]],
        ["取引先コード", order["customerCode"], "取引先名", order["customerName"]],
        ["納品希望日", order["deliveryDate"], "支払条件", order["paymentTerms"]],
        ["商品コード", order["productCode"], "品名", order["productName"]],
        ["数量", f"{order['quantity']:,}", "単価", f"{order['unitPrice']:,} {order['currency']}"],
        ["出荷方法", order["shippingMethod"], "検収要否", order["inspection"]],
        ["発注担当者", order["picName"], "通貨", order["currency"]],
    ]
    story = [
        Paragraph("注　文　書", styles["title"]),
        Spacer(1, 6 * mm),
        Paragraph("株式会社サンプル商事 御中", styles["normal"]),
        Spacer(1, 4 * mm),
        _table(rows, [30 * mm, 45 * mm, 30 * mm, 55 * mm]),
        Spacer(1, 6 * mm),
        Paragraph("備考: " + order["remarks"], styles["normal"]),
        Spacer(1, 10 * mm),
        Paragraph("本注文書はハンズオン用のサンプルです。実在の企業・取引とは関係ありません。", styles["small"]),
    ]
    doc.build(story)


def build_pdf_format_b(order, path):
    """フォーマット B: ヘッダー情報 + 明細行テーブル型。"""
    styles = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm)
    header_rows = [
        ["ORDER NO.", order["orderNumber"]],
        ["ISSUE DATE (受注日)", order["orderDate"]],
        ["SUPPLIER (取引先)", f"{order['customerName']} ({order['customerCode']})"],
        ["REQUESTED DELIVERY (納品希望日)", order["deliveryDate"]],
        ["BUYER (担当者)", order["picName"]],
    ]
    detail_rows = [
        ["ITEM CODE", "DESCRIPTION", "QTY", "UNIT PRICE", "CURRENCY", "AMOUNT"],
        [
            order["productCode"],
            order["productName"],
            f"{order['quantity']:,}",
            f"{order['unitPrice']:,}",
            order["currency"],
            f"{order['quantity'] * order['unitPrice']:,}",
        ],
    ]
    terms_rows = [
        ["PAYMENT TERMS (支払条件)", order["paymentTerms"]],
        ["SHIPPING (出荷方法)", order["shippingMethod"]],
        ["INSPECTION (検収要否)", order["inspection"]],
        ["REMARKS (備考)", order["remarks"]],
    ]
    story = [
        Paragraph("PURCHASE ORDER / 発注依頼書", styles["title"]),
        Spacer(1, 5 * mm),
        _table(header_rows, [65 * mm, 100 * mm]),
        Spacer(1, 6 * mm),
        Paragraph("■ 明細", styles["normal"]),
        Spacer(1, 2 * mm),
        _table(detail_rows, [25 * mm, 55 * mm, 20 * mm, 27 * mm, 20 * mm, 28 * mm], header_rows=1),
        Spacer(1, 6 * mm),
        Paragraph("■ 取引条件", styles["normal"]),
        Spacer(1, 2 * mm),
        _table(terms_rows, [55 * mm, 110 * mm]),
        Spacer(1, 8 * mm),
        Paragraph("This document is a sample for the Playwright hands-on.", styles["small"]),
    ]
    doc.build(story)


def build_pdf_format_c(order, path):
    """フォーマット C: FAX 送信票風の縦並びレイアウト。"""
    styles = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm)
    lines = [
        ("件名", "受注内容確認のご連絡"),
        ("受注番号", order["orderNumber"]),
        ("受注年月日", order["orderDate"]),
        ("お得意先", f"{order['customerCode']}　{order['customerName']} 様"),
        ("ご希望納期", order["deliveryDate"]),
        ("ご注文品番", order["productCode"]),
        ("品名", order["productName"]),
        ("ご注文数量", f"{order['quantity']:,} 個"),
        ("販売単価", f"{order['unitPrice']:,} 円 ({order['currency']})"),
        ("お支払条件", order["paymentTerms"]),
        ("配送区分", order["shippingMethod"]),
        ("検収", order["inspection"]),
        ("弊社担当", order["picName"]),
        ("連絡事項", order["remarks"]),
    ]
    story = [
        Paragraph("ＦＡＸ送信票 （受注確認）", styles["title"]),
        Spacer(1, 4 * mm),
        Paragraph("送信日: " + order["orderDate"] + "　/　送信枚数: 1 枚", styles["normal"]),
        Spacer(1, 6 * mm),
    ]
    for label, value in lines:
        story.append(Paragraph(f"{label}　：　{value}", styles["normal"]))
        story.append(Spacer(1, 1.5 * mm))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("※ 本 FAX はハンズオン用のサンプルです。", styles["small"]))
    doc.build(story)


HEADER_FILL = PatternFill("solid", fgColor="DDE5EE")
TITLE_FONT = Font(size=14, bold=True)
LABEL_FONT = Font(bold=True)
THIN = Side(style="thin", color="8C99A8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_pairs(ws, start_row, pairs, label_col="A", value_col="B"):
    row = start_row
    for label, value in pairs:
        ws[f"{label_col}{row}"] = label
        ws[f"{label_col}{row}"].font = LABEL_FONT
        ws[f"{label_col}{row}"].fill = HEADER_FILL
        ws[f"{label_col}{row}"].border = BORDER
        ws[f"{value_col}{row}"] = value
        ws[f"{value_col}{row}"].border = BORDER
        row += 1
    return row


def build_excel_format_a(order, path):
    """フォーマット A: 1 シート縦並びのオーソドックスな注文書。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "注文書"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 38
    ws["A1"] = "注文書"
    ws["A1"].font = TITLE_FONT
    _write_pairs(
        ws,
        3,
        [
            ("注文番号", order["orderNumber"]),
            ("受注日", order["orderDate"]),
            ("取引先コード", order["customerCode"]),
            ("取引先名", order["customerName"]),
            ("納品希望日", order["deliveryDate"]),
            ("商品コード", order["productCode"]),
            ("品名", order["productName"]),
            ("数量", order["quantity"]),
            ("単価", order["unitPrice"]),
            ("通貨", order["currency"]),
            ("支払条件", order["paymentTerms"]),
            ("出荷方法", order["shippingMethod"]),
            ("検収要否", order["inspection"]),
            ("担当者名", order["picName"]),
            ("備考", order["remarks"]),
        ],
    )
    wb.save(path)


def build_excel_format_b(order, path):
    """フォーマット B: ヘッダーシートと明細シートに分かれた形式。"""
    wb = Workbook()
    header = wb.active
    header.title = "ヘッダ"
    header.column_dimensions["A"].width = 24
    header.column_dimensions["B"].width = 38
    header["A1"] = "受発注ヘッダ情報"
    header["A1"].font = TITLE_FONT
    _write_pairs(
        header,
        3,
        [
            ("Order No.", order["orderNumber"]),
            ("Order Date", order["orderDate"]),
            ("Customer Code", order["customerCode"]),
            ("Customer Name", order["customerName"]),
            ("Requested Delivery Date", order["deliveryDate"]),
            ("Payment Terms", order["paymentTerms"]),
            ("Shipping Method", order["shippingMethod"]),
            ("Person In Charge", order["picName"]),
            ("Inspection", order["inspection"]),
            ("Remarks", order["remarks"]),
        ],
    )

    detail = wb.create_sheet("明細")
    headers = ["商品コード", "品名", "数量", "単価", "通貨", "金額"]
    for index, value in enumerate(headers, start=1):
        cell = detail.cell(row=1, column=index, value=value)
        cell.font = LABEL_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
        detail.column_dimensions[cell.column_letter].width = 18
    values = [
        order["productCode"],
        order["productName"],
        order["quantity"],
        order["unitPrice"],
        order["currency"],
        order["quantity"] * order["unitPrice"],
    ]
    for index, value in enumerate(values, start=1):
        detail.cell(row=2, column=index, value=value).border = BORDER

    wb.save(path)


def build_excel_format_c(order, path):
    """フォーマット C: 受注一覧表 (1 行 1 受注) の横持ち形式。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "受注一覧"
    ws["A1"] = "受注一覧表 (自社システム出力)"
    ws["A1"].font = TITLE_FONT

    headers = [
        "受注NO",
        "受注日",
        "得意先CD",
        "得意先名",
        "希望納期",
        "品目CD",
        "品目名",
        "受注数",
        "売単価",
        "通貨CD",
        "支払条件",
        "配送区分",
        "検収",
        "担当",
        "摘要",
    ]
    values = [
        order["orderNumber"],
        order["orderDate"],
        order["customerCode"],
        order["customerName"],
        order["deliveryDate"],
        order["productCode"],
        order["productName"],
        order["quantity"],
        order["unitPrice"],
        order["currency"],
        order["paymentTerms"],
        order["shippingMethod"],
        order["inspection"],
        order["picName"],
        order["remarks"],
    ]
    for index, value in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=index, value=value)
        cell.font = LABEL_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(12, len(str(value)) + 8)
    for index, value in enumerate(values, start=1):
        ws.cell(row=4, column=index, value=value).border = BORDER

    wb.save(path)


EXCEL_ORDERS = [
    {
        "file": "order-101",
        "orderNumber": "PO-2026-0208",
        "orderDate": "2026/05/18",
        "customerCode": "C-1002",
        "customerName": "東京電子工業",
        "deliveryDate": "2026/06/05",
        "productCode": "P-2003",
        "productName": "樹脂カバー TypeB",
        "quantity": 500,
        "unitPrice": 640,
        "currency": "JPY",
        "paymentTerms": "60日後支払",
        "shippingMethod": "通常便",
        "picName": "高橋 美咲",
        "inspection": "不要",
        "remarks": "色番は前回同様 RAL7035 とする。",
    },
    {
        "file": "order-102",
        "orderNumber": "PO-2026-0233",
        "orderDate": "2026/06/02",
        "customerCode": "C-1004",
        "customerName": "関西精密機器",
        "deliveryDate": "2026/07/10",
        "productCode": "P-2004",
        "productName": "ベアリングユニット",
        "quantity": 80,
        "unitPrice": 5400,
        "currency": "JPY",
        "paymentTerms": "30日後支払",
        "shippingMethod": "急送便",
        "picName": "伊藤 大輔",
        "inspection": "必要",
        "remarks": "梱包は 10 個単位。",
    },
    {
        "file": "order-103",
        "orderNumber": "PO-2026-0290",
        "orderDate": "2026/06/22",
        "customerCode": "C-1001",
        "customerName": "山田製作所",
        "deliveryDate": "2026/08/03",
        "productCode": "P-2002",
        "productName": "ステンレス板 SUS304",
        "quantity": 24,
        "unitPrice": 13200,
        "currency": "USD",
        "paymentTerms": "前払い",
        "shippingMethod": "引取",
        "picName": "渡辺 里奈",
        "inspection": "不要",
        "remarks": "海外向け案件のため通貨は USD。",
    },
]

PDF_BUILDERS = [build_pdf_format_a, build_pdf_format_b, build_pdf_format_c]
EXCEL_BUILDERS = [build_excel_format_a, build_excel_format_b, build_excel_format_c]


def main() -> None:
    pdfmetrics.registerFont(UnicodeCIDFont(FONT_NAME))
    os.makedirs(PDF_DIR, exist_ok=True)
    os.makedirs(EXCEL_DIR, exist_ok=True)

    for order, builder in zip(PDF_ORDERS, PDF_BUILDERS):
        pdf_path = os.path.join(PDF_DIR, f"{order['file']}.pdf")
        builder(order, pdf_path)
        print(f"generated: {pdf_path}")

    for order, builder in zip(EXCEL_ORDERS, EXCEL_BUILDERS):
        excel_path = os.path.join(EXCEL_DIR, f"{order['file']}.xlsx")
        builder(order, excel_path)
        print(f"generated: {excel_path}")


if __name__ == "__main__":
    main()
